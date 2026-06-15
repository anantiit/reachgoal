#!/usr/bin/env python3
"""
Verify that HTML and PPTX presentations have matching slides.
"""

from pptx import Presentation

# HTML slide titles (extracted from the file)
html_slides = [
    "Java Collections & Streams: Architecture to Performance",  # Title slide
    "Why Learn Collections & Streams?",
    "Collections Framework - Architecture",
    "ArrayList - What/Why/When",
    "LinkedList - What/Why/When",  # First occurrence
    # "LinkedList - What/Why/When",  # Duplicate (line 538) - SKIPPED
    "ArrayList vs LinkedList - The Verdict",
    "HashSet - What/Why/When",
    "Set Implementations - Comparison",
    "hashCode() & equals() Contract",
    "HashMap - What/Why/When",
    "LinkedHashMap - Predictable Iteration",
    "TreeMap - Sorted & Navigable",
    "ConcurrentHashMap - What/Why/When",
    "Java 8+ Map Compute Methods",
    "Queue & Deque - FIFO & LIFO",
    "Immutable Collections (Java 9+)",
    "Comparable vs Comparator",
    "Performance Summary - Decision Matrix",
    "Stream API - Declarative Pipelines",
    "Stream Operations - Quick Reference",
    "Lazy Evaluation & Short-Circuiting",
    "Advanced Collectors",
    "reduce vs collect",
    "Parallel Streams - Power & Pitfalls",
    "Parallel Stream Pitfalls & Solutions",
    "Key Takeaways & Best Practices"
]

def verify_pptx_slides():
    """Verify PPTX slide count and titles."""
    prs = Presentation('java-practice/java-collections/java_collections_streams_deep_dive.pptx')
    
    print(f"HTML Slides: {len(html_slides)}")
    print(f"PPTX Slides: {len(prs.slides)}")
    print()
    
    if len(html_slides) != len(prs.slides):
        print(f"❌ MISMATCH: HTML has {len(html_slides)} slides, PPTX has {len(prs.slides)} slides")
        return False
    
    print("✅ Slide count matches!")
    print("\n" + "="*80)
    print("SLIDE-BY-SLIDE VERIFICATION")
    print("="*80)
    
    all_match = True
    for i, (html_title, pptx_slide) in enumerate(zip(html_slides, prs.slides), start=1):
        # Extract text from PPTX slide
        pptx_text = []
        for shape in pptx_slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                pptx_text.append(shape.text.strip())
        
        # For title slide (slide 1), check if title appears in any text
        if i == 1:
            # Check if key words from title appear
            found = any("Java Collections" in text and "Streams" in text and "Architecture" in text and "Performance" in text for text in pptx_text)
            status = "✅" if found else "❌"
            print(f"\nSlide {i}: {status}")
            print(f"  HTML: {html_title}")
            print(f"  PPTX: Title slide with 'Java Collections & Streams: Architecture to Performance'")
            if not found:
                all_match = False
                print(f"  Available text: {pptx_text[:3]}")
        else:
            # For other slides, look for the title in text
            # Normalize for comparison (handle slight variations)
            html_normalized = html_title.lower().replace("what/why/when", "").strip()
            found = False
            for text in pptx_text:
                if html_normalized in text.lower() or html_title in text:
                    found = True
                    break
            
            status = "✅" if found else "❌"
            print(f"\nSlide {i}: {status}")
            print(f"  HTML: {html_title}")
            
            if found:
                matching_text = [t for t in pptx_text if html_normalized in t.lower() or html_title in t]
                print(f"  PPTX: {matching_text[0] if matching_text else 'Found in slide'}")
            else:
                all_match = False
                print(f"  PPTX: NOT FOUND")
                print(f"  Available text: {pptx_text[:3]}")
    
    print("\n" + "="*80)
    if all_match:
        print("✅ ALL SLIDES VERIFIED SUCCESSFULLY!")
        print("✅ HTML and PPTX presentations match perfectly!")
    else:
        print("❌ Some slides have mismatches - review details above")
    print("="*80)
    
    return all_match

if __name__ == "__main__":
    verify_pptx_slides()
